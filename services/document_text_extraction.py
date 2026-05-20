"""Ephemeral document text extraction for assistant uploads."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from backend_common import compact_text


MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_TEXT_CHARS = 50_000
TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".csv"}
PDF_CONTENT_TYPES = {"application/pdf"}
DOCX_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
}


class UnsupportedDocumentError(ValueError):
    """Raised when an uploaded file cannot be converted into text."""


class DocumentExtractionError(ValueError):
    """Raised when a supported file cannot be parsed."""


@dataclass(frozen=True)
class ExtractedDocument:
    name: str
    text: str
    truncated: bool


def extract_document_text(
    filename: str,
    content_type: str | None,
    payload: bytes,
) -> ExtractedDocument:
    """Extract bounded text from a supported, client-provided document."""
    safe_name = Path(filename or "document").name
    extension = Path(safe_name).suffix.lower()
    media_type = (content_type or "").split(";")[0].strip().lower()

    if len(payload) > MAX_UPLOAD_BYTES:
        raise UnsupportedDocumentError("File is larger than the 10 MB upload limit.")

    if extension in TEXT_EXTENSIONS or media_type.startswith("text/"):
        text = _extract_plain_text(payload)
    elif extension == ".pdf" or media_type in PDF_CONTENT_TYPES:
        text = _extract_pdf_text(payload)
    elif extension == ".docx" or media_type in DOCX_CONTENT_TYPES:
        text = _extract_docx_text(payload)
    else:
        raise UnsupportedDocumentError(
            "Supported uploads are PDF, DOCX, TXT, Markdown, and CSV files."
        )

    cleaned = _clean_extracted_text(text)
    if not cleaned:
        raise DocumentExtractionError("No readable text could be extracted from this file.")

    truncated = len(cleaned) > MAX_EXTRACTED_TEXT_CHARS
    return ExtractedDocument(
        name=safe_name,
        text=cleaned[:MAX_EXTRACTED_TEXT_CHARS],
        truncated=truncated,
    )


def _extract_plain_text(payload: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentExtractionError("Text file encoding is not supported.")


def _extract_pdf_text(payload: bytes) -> str:
    try:
        import fitz

        with fitz.open(stream=payload, filetype="pdf") as document:
            return "\n\n".join(page.get_text("text") for page in document)
    except Exception as exc:  # pragma: no cover - parser internals vary.
        raise DocumentExtractionError("PDF text could not be extracted.") from exc


def _extract_docx_text(payload: bytes) -> str:
    try:
        import docx

        document = docx.Document(BytesIO(payload))
    except Exception as exc:  # pragma: no cover - parser internals vary.
        raise DocumentExtractionError("DOCX text could not be extracted.") from exc

    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join([*paragraphs, *table_cells])


def _clean_extracted_text(text: str) -> str:
    lines = [compact_text(line) for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()
